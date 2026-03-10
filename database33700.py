import os
import sqlite3
from pandas import DataFrame
from datetime import datetime
from openpyxl import Workbook

from convertisseur import construire_ordre_colonnes_disponibles


def exporter_vers_db(
        db_path: str,
        dfs: list[DataFrame],
        noms_fichier: list[str],
) -> dict:
    """
    Pour chaque (df, nom_fichier), insère df dans base_donnees seulement si
    nom_fichier n'est pas déjà présent dans fichiers_charges.
    Met à jour fichiers_charges pour les nouveaux fichiers.

    Retourne un bilan.
    """
    if len(dfs) != len(noms_fichier):
        raise ValueError("dfs et noms_fichier doivent avoir la même longueur (même ordre).")

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    bilan = {
        "fichiers_inseres": [],
        "fichiers_ignores_deja_charges": [],
        "fichiers_ignores_df_vide": [],
        "lignes_inserees_total": 0,
        "lignes_inserees_par_fichier": {},  # nom_fichier -> nb lignes
    }

    with sqlite3.connect(db_path) as conn:
        cur = conn.cursor()

        # Tables
        cur.execute("""
        CREATE TABLE IF NOT EXISTS base_donnees (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            DATE_SIGNALEMENT TEXT,
            MESSAGE TEXT,
            EMETTEUR TEXT,
            ALIAS_SIGNALANT TEXT,
            NUMERO_REBOND_SIGNAL TEXT,
            OPERATEUR_SIGNALANT TEXT,
            URL_REBOND_SIGNALE TEXT,
            date_requalifiee TEXT,
            expediteur_nettoye TEXT,
            typologie_expediteur TEXT,
            operateur_arcep TEXT,
            typologie_rebond TEXT,
            categorie_no_cible TEXT,
            mois TEXT,
            DATE_RECEPTION TEXT,
            MOIS_RECEPTION TEXT,
            ANALYSE_STOP TEXT,
            TYPE_EMETTEUR TEXT,
            CANAL TEXT,
            EMETTEUR_NETTOYE TEXT,
            rebond_nettoye TEXT,
            type_protection TEXT,
            phishing TEXT,
            mots_clefs TEXT,
            score_smishing TEXT,
            tous_les_mots_clefs TEXT,
            opr_arcep_rebond TEXT,
            traitements TEXT
        )
        """)

        cur.execute("""
        CREATE TABLE IF NOT EXISTS fichiers_charges (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nom_fichier TEXT UNIQUE,
            date_charge TEXT
        )
        """)

        # Charger en mémoire l’ensemble des fichiers déjà chargés (rapide)
        deja_charges = {r[0] for r in cur.execute("SELECT nom_fichier FROM fichiers_charges").fetchall()}

        # Colonnes attendues (sans id)
        cols_db = [r[1] for r in cur.execute("PRAGMA table_info(base_donnees)").fetchall()]
        cols_db = [c for c in cols_db if c != "id"]

        for df, nom in zip(dfs, noms_fichier):
            # 1) déjà chargé ?
            if nom in deja_charges:
                bilan["fichiers_ignores_deja_charges"].append(nom)
                continue

            # 2) df vide ?
            if df is None or df.empty:
                bilan["fichiers_ignores_df_vide"].append(nom)
                # On peut choisir de marquer quand même comme "chargé" ou pas.
                # Ici: on NE le marque PAS.
                continue

            df_a_inserer = df.copy()

            # # 3) (Optionnel) normaliser dates si présentes
            # for col in ["dt_signalement", "dt_fraude"]:
            #     if col in df_a_inserer.columns:
            #         df_a_inserer[col] = pd.to_datetime(df_a_inserer[col], errors="coerce").dt.strftime("%Y-%m-%d %H:%M:%S")

            # 4) aligner les colonnes sur la table
            # → colonnes manquantes = NaN, colonnes en trop = ignorées
            df_a_inserer = df_a_inserer.reindex(columns=cols_db)

            # 5) insérer en DB
            df_a_inserer.to_sql("base_donnees", conn, if_exists="append", index=False)

            # 6) mettre à jour fichiers_charges
            cur.execute(
                "INSERT INTO fichiers_charges (nom_fichier, date_charge) VALUES (?, ?)",
                (nom, now)
            )
            deja_charges.add(nom)

            nb = int(len(df_a_inserer))
            bilan["fichiers_inseres"].append(nom)
            bilan["lignes_inserees_par_fichier"][nom] = nb
            bilan["lignes_inserees_total"] += nb

    return bilan


# def exporter_base_vers_excel(db_path, outdir, filepath, max_rows=1_000_000):
#     table = "base_donnees"
#     fetch_batch = 10000
#
#     conn = sqlite3.connect(db_path)
#     cur = conn.cursor()
#
#     # 1️⃣ compter les lignes par mois
#     cur.execute(f"""
#         SELECT mois, COUNT(*) AS n
#         FROM {table}
#         WHERE mois IS NOT NULL AND TRIM(mois) <> ''
#         GROUP BY mois
#         ORDER BY mois DESC
#     """)
#
#     rows = cur.fetchall()
#
#     selected_months = []
#     total = 0
#
#     for mois, n in rows:
#         if total + n > max_rows:
#             break
#         selected_months.append(mois)
#         total += n
#
#     if not selected_months:
#         print("Aucun mois sélectionné.")
#         conn.close()
#         return
#
#     print("Mois retenus :", selected_months)
#     print("Total lignes :", total)
#
#     # 2️⃣ requête pour récupérer les données
#     placeholders = ",".join(["?"] * len(selected_months))
#
#     query = f"""
#         SELECT *
#         FROM {table}
#         WHERE mois IN ({placeholders})
#         ORDER BY mois DESC, id ASC
#     """
#
#     cur.execute(query, selected_months)
#
#     # 3️⃣ création Excel (mode streaming)
#     wb = Workbook(write_only=True)
#     ws = wb.create_sheet("export")
#
#     # en-têtes
#     columns = [d[0] for d in cur.description]
#     ws.append(columns)
#
#     # écriture par batch
#     while True:
#         batch = cur.fetchmany(fetch_batch)
#         if not batch:
#             break
#         for row in batch:
#             ws.append(row)
#     output_path = os.path.join(outdir, os.path.basename(filepath).split('.')[0] + '.xlsx')
#     wb.save(output_path)
#
#     conn.close()
#
#     print(f"Export terminé → {output_path}")

def exporter_base_vers_excel(db_path, outdir, filepath, max_rows=1_000_000):
    table = "base_donnees"
    fetch_batch = 10000

    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    # 1) compter les lignes par mois
    cur.execute(f"""
        SELECT mois, COUNT(*) AS n
        FROM {table}
        WHERE mois IS NOT NULL AND TRIM(mois) <> ''
        GROUP BY mois
        ORDER BY mois DESC
    """)

    rows = cur.fetchall()

    selected_months = []
    total = 0

    for mois, n in rows:
        if total + n > max_rows:
            break
        selected_months.append(mois)
        total += n

    if not selected_months:
        print("Aucun mois sélectionné.")
        conn.close()
        return

    print("Mois retenus :", selected_months)
    print("Total lignes :", total)

    # 2) récupérer la liste réelle des colonnes de la table
    cur.execute(f"PRAGMA table_info({table})")
    table_columns = [row[1] for row in cur.fetchall()]  # row[1] = nom de colonne

    ordered_columns = construire_ordre_colonnes_disponibles(table_columns)

    placeholders = ",".join(["?"] * len(selected_months))
    cols_sql = ", ".join(f'"{col}"' for col in ordered_columns)

    query = f"""
        SELECT {cols_sql}
        FROM {table}
        WHERE mois IN ({placeholders})
        ORDER BY mois DESC, id ASC
    """

    cur.execute(query, selected_months)

    # 3) création Excel (mode streaming)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("export")

    # en-têtes dans le bon ordre
    ws.append(ordered_columns)

    # écriture par batch
    while True:
        batch = cur.fetchmany(fetch_batch)
        if not batch:
            break
        for row in batch:
            ws.append(row)

    output_path = os.path.join(
        outdir,
        os.path.basename(filepath).split('.')[0] + '.xlsx'
    )
    wb.save(output_path)

    conn.close()

    print(f"Export terminé → {output_path}")
